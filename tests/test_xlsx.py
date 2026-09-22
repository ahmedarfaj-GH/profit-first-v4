import zipfile
from io import BytesIO

import openpyxl
import pytest

from app.engine.xlsx_template import TemplateParseError, build_template_workbook, parse_intake_workbook

VALUES_BY_SERIAL = {
    2: "01/07/2026", 3: "31/08/2026", 6: 6850.37, 7: 119142.98, 8: 59231.64,
    9: 0, 10: 41231.64, 11: 11384.9, 12: 0, 14: 20000,
}


def filled_template(overrides=None):
    wb = openpyxl.load_workbook(build_template_workbook("محل آيس كريم"))
    ws = wb.active
    values = {**VALUES_BY_SERIAL, **(overrides or {})}
    for row in ws.iter_rows(min_row=4):
        serial = row[0].value
        if serial in values:
            row[4].value = values[serial]
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_generated_template_round_trips():
    parsed = parse_intake_workbook(filled_template())
    assert parsed["period_start"] == "2026-07-01" and parsed["period_end"] == "2026-08-31"
    assert parsed["inputs"]["opening_cash_balance"] == 6850.37
    assert parsed["inputs"]["payroll_due"] == 41231.64
    assert parsed["inputs"]["operational_reserve_target"] == 20000
    assert parsed["inputs"]["royalty_due"] == 0


def test_parser_tolerates_extra_columns_and_shifted_header():
    """Real files sometimes insert due-date/priority columns and start lower on the sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["title"] * 3)
    ws.append([])
    ws.append([])
    ws.append(["الرقم التسلسلي", "مجموعة", "اسم البيان المطلوب", "وصف", "القيمة", "نوع", "حالة", "الفترة",
               "تاريخ الاستحقاق", "الأولوية", "الهدف الكامل", "المرجع الداخلي", "ملاحظات"])
    for serial, value in VALUES_BY_SERIAL.items():
        ws.append([serial, "g", f"item {serial}", "d", value, "t", "s", "p", "30/08/2026", 1, 20000, "ref", "n"])
    buf = BytesIO()
    wb.save(buf)
    parsed = parse_intake_workbook(buf.getvalue())
    assert parsed["inputs"]["total_collections"] == 119142.98
    assert parsed["period_end"] == "2026-08-31"


def test_negative_values_are_rejected():
    with pytest.raises(TemplateParseError):
        parse_intake_workbook(filled_template({7: -5}))


def test_missing_required_rows_are_reported():
    wb = openpyxl.load_workbook(build_template_workbook())
    ws = wb.active
    for row in list(ws.iter_rows(min_row=4)):
        if row[0].value == 14:
            ws.delete_rows(row[0].row)
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(TemplateParseError, match="ناقصة"):
        parse_intake_workbook(buf.getvalue())


def test_non_excel_bytes_are_rejected():
    with pytest.raises(TemplateParseError):
        parse_intake_workbook(b"this is not a spreadsheet")


def test_zip_bomb_is_rejected_before_parsing():
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/huge.xml", b"0" * (25 * 1024 * 1024))
    assert len(buf.getvalue()) < 100_000
    with pytest.raises(TemplateParseError, match="أكبر"):
        parse_intake_workbook(buf.getvalue())
